"""Build an Excel workbook for one target table's lineage summary."""
import io
import sqlite3
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

_HEADER_FILL = PatternFill("solid", fgColor="1F4E78")
_TARGET_FILL = PatternFill("solid", fgColor="DBEAFE")
_SOURCE_FILL = PatternFill("solid", fgColor="CFFAFE")
_HEADER_FONT = Font(bold=True, color="FFFFFF")
_WRAP = Alignment(vertical="top", wrap_text=True)

# Keep these in step with the UI badges (src/ui/src/components/relationStyle.ts).
RELATION_LABEL = {
    "EXPLICIT_INATTR": "Direct mapping",
    "EXPLICIT_PATHFROM_PATHTO": "Direct mapping (path)",
    "AUTO_MAP_BY_NAME": "Auto-mapped by name",
    "TRANSFORMATION_AUTO_MAP_BY_NAME": "Transformation auto-map",
    "EXPRESSION": "Expression",
    "CONSTANT_ASSIGNMENT": "Constant",
    "ENRICHMENT_LOOKUP": "Lookup",
    "RUNTIME_FUNCTION_ASSIGNMENT": "Function",
    "NOT_PERSISTED": "Not persisted",
}

KIND_LABEL = {"column": "Column", "constant": "Constant", "function": "Function", "not_persisted": "Not persisted"}

_KIND_MEANING = {
    "Column": "The lineage ends at a column of a source table or file.",
    "Constant": "The column is filled by a constant value; nothing is read from a source.",
    "Function": "The column is filled by a runtime function; nothing is read from a source.",
    "Not persisted": "The value is computed and not stored; it is not read from a source.",
}


def _label(relation_type: str) -> str:
    return RELATION_LABEL.get(relation_type, relation_type.replace("_", " ").capitalize())


def _hops(source: dict[str, Any]) -> str:
    lo, hi = source["hops_min"], source["hops_max"]
    return f"{lo} hop" if lo == hi == 1 else (str(lo) if lo == hi else f"{lo}–{hi} hops")


def _extraction_timestamp(conn: sqlite3.Connection) -> str:
    try:
        row = conn.execute("SELECT MAX(extraction_timestamp) FROM extraction_run").fetchone()
        return row[0] or "unknown"
    except sqlite3.Error:
        return "unknown"


def _db_name(conn: sqlite3.Connection) -> str:
    try:
        return str(conn.execute("PRAGMA database_list").fetchone()[2]).rsplit("/", 1)[-1]
    except (sqlite3.Error, TypeError):
        return "unknown"


def _style_sheet(sheet, widths: list[int], wrap_from_row: int = 2) -> None:
    for cell in sheet[1]:
        cell.font = _HEADER_FONT
        cell.fill = _HEADER_FILL
        cell.alignment = _WRAP
    sheet.freeze_panes = "A2"
    for i, width in enumerate(widths, start=1):
        sheet.column_dimensions[get_column_letter(i)].width = width
    for row in sheet.iter_rows(min_row=wrap_from_row):
        for cell in row:
            cell.alignment = _WRAP


def _summary_sheet(wb: Workbook, object_name: str, columns: list[dict[str, Any]]) -> None:
    """One row per target column and ultimate source: where the value finally comes from."""
    sheet = wb.active
    sheet.title = "Summary"
    sheet.append(["Target table", "Target column", "Ultimate source", "Source type", "Hops", "Path variants", "EXPR (driving)", "OUTPUT (result)"])
    for column in columns:
        sources = column["sources"] or [None]
        for source in sources:
            if source is None:
                sheet.append([object_name, column["column_name"], "No lineage", "", "", "", "", ""])
            else:
                sheet.append([
                    object_name,
                    column["column_name"],
                    f"{source['object_name']}.{source['column_name']}",
                    KIND_LABEL.get(source["kind"], source["kind"]),
                    _hops(source),
                    source["route_count"],
                    source["condition"] or "",
                    source["result_expression"] or "",
                ])
            r = sheet.max_row
            sheet.cell(row=r, column=2).fill = _TARGET_FILL
            sheet.cell(row=r, column=3).fill = _SOURCE_FILL
    _style_sheet(sheet, [26, 24, 36, 14, 12, 12, 40, 40])


def _steps_sheet(wb: Workbook, object_name: str, columns: list[dict[str, Any]]) -> None:
    """One row per hop: target column -> hop 1 -> hop 2 -> ... -> ultimate source."""
    sheet = wb.create_sheet("Steps")
    sheet.append(["Target table", "Target column", "Ultimate source", "Hop", "From", "To", "Mapping kind", "Component", "EXPR (driving)", "OUTPUT (result)"])
    for column in columns:
        for source in column["sources"]:
            route = source["route"]
            terminal = f"{source['object_name']}.{source['column_name']}"
            for index, step in enumerate(route):
                is_leaf = index == len(route) - 1
                sheet.append([
                    object_name,
                    column["column_name"],
                    terminal,
                    index + 1,
                    step["endpoint"],
                    route[index - 1]["endpoint"] if index else f"{object_name}.{column['column_name']}",
                    _label(step["relation_type"]),
                    step["component_name"] or "",
                    (source["condition"] or "") if is_leaf else "",
                    (source["result_expression"] or "") if is_leaf else "",
                ])
                r = sheet.max_row
                if index == 0:
                    sheet.cell(row=r, column=6).fill = _TARGET_FILL
                if is_leaf:
                    sheet.cell(row=r, column=5).fill = _SOURCE_FILL
    _style_sheet(sheet, [26, 24, 36, 6, 36, 36, 22, 24, 40, 40])


def _legend_sheet(wb: Workbook, conn: sqlite3.Connection, project: str) -> None:
    sheet = wb.create_sheet("How to read")
    rows = [
        ("Project", project),
        ("Database", _db_name(conn)),
        ("Extracted at", _extraction_timestamp(conn)),
        ("", ""),
        ("Summary sheet", "One row per target column and ultimate source. 'Hops' counts the mappings between them; 'Path variants' counts distinct paths when there is more than one."),
        ("Steps sheet", "One row per hop of the shortest path from the target column to each ultimate source. 'From' is the column being read; 'To' is what it feeds. When Path variants > 1, only one path is listed."),
        ("", ""),
        ("Mapping kinds", ""),
        *[(_label(k), "") for k in RELATION_LABEL],
        ("", ""),
        ("Source types", ""),
        *[(label, _KIND_MEANING[label]) for label in _KIND_MEANING],
    ]
    for pair in rows:
        sheet.append(list(pair))
    sheet.column_dimensions["A"].width = 34
    sheet.column_dimensions["B"].width = 110
    for row in sheet.iter_rows():
        for cell in row:
            cell.alignment = _WRAP


def build_table_workbook(conn: sqlite3.Connection, project: str, object_name: str, summary: dict[str, Any]) -> bytes:
    wb = Workbook()
    _summary_sheet(wb, object_name, summary["columns"])
    _steps_sheet(wb, object_name, summary["columns"])
    _legend_sheet(wb, conn, project)
    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()
