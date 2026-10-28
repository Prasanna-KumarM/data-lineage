import io
import sqlite3
from pathlib import Path

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from conftest import ROOT

import sys
sys.path.insert(0, str(ROOT / "src" / "api"))

from lineage_api.app import create_app  # noqa: E402


@pytest.fixture()
def sp_clo_client():
    path = ROOT / "outputs" / "SP_CLO" / "lineage.db"
    if not path.is_file():
        pytest.skip("SP_CLO lineage DB is not available")
    return TestClient(create_app(str(path)))


def _book(response):
    assert response.status_code == 200
    assert response.headers["content-type"].startswith("application/vnd.openxmlformats")
    assert 'filename="SP_CLO_CM_PAR_LOAN_TRADES_lineage.xlsx"' in response.headers["content-disposition"]
    return load_workbook(io.BytesIO(response.content))


def _target_id(client, name):
    targets = client.get("/targets", params={"project": "SP_CLO"}).json()
    return next(t["object_id"] for t in targets if t["object_name"] == name)


def test_workbook_rows_match_the_summary(sp_clo_client):
    object_id = _target_id(sp_clo_client, "CM_PAR_LOAN_TRADES")
    summary = sp_clo_client.get(f"/targets/{object_id}/summary", params={"project": "SP_CLO"}).json()
    wb = _book(sp_clo_client.get(f"/targets/{object_id}/export.xlsx", params={"project": "SP_CLO"}))
    assert wb.sheetnames == ["Summary", "Steps", "How to read"]

    sheet = wb["Summary"]
    headers = [c.value for c in sheet[1]]
    assert headers == ["Target table", "Target column", "Ultimate source", "Source type", "Hops", "Path variants", "When (condition)", "Value (result)"]
    rows = [[c.value for c in row] for row in sheet.iter_rows(min_row=2)]
    assert len(rows) == sum(max(1, len(c["sources"])) for c in summary["columns"])
    assert len({r[0] for r in rows}) == 1
    assert sorted({r[1] for r in rows}) == sorted(c["column_name"] for c in summary["columns"])

    fdk = [r for r in rows if r[1] == "FDK_MT"]
    assert len(fdk) == len(next(c for c in summary["columns"] if c["column_name"] == "FDK_MT")["sources"]) == 3
    assert {r[2] for r in fdk} == {
        "CM_CLO_OFFER.OFR_HTCS_BOOK_CODE",
        "CM_CLO_OFFER.OFR_HTC_BOOK_CODE",
        "CM_CLO_OFFER.OFR_HTS_BOOK_CODE",
    }
    assert all(r[3] == "Column" for r in fdk)

    steps = [[c.value for c in row] for row in wb["Steps"].iter_rows(min_row=2)]
    assert [c.value for c in wb["Steps"][1]] == ["Target table", "Target column", "Ultimate source", "Hop", "From", "To", "Mapping kind", "Component", "When (condition)", "Value (result)"]
    assert len(steps) == sum(len(s["route"]) for c in summary["columns"] for s in c["sources"])
    fdk_steps = [r for r in steps if r[1] == "FDK_MT" and r[2] == "CM_CLO_OFFER.OFR_HTCS_BOOK_CODE"]
    assert [r[3] for r in fdk_steps] == list(range(1, len(fdk_steps) + 1))
    assert fdk_steps[0][5] == "CM_PAR_LOAN_TRADES.FDK_MT"
    assert fdk_steps[-1][4] == "CM_CLO_OFFER.OFR_HTCS_BOOK_CODE"
    for step, nxt in zip(fdk_steps, fdk_steps[1:]):
        assert step[4] == nxt[5]


def test_columns_without_lineage_appear_once_with_a_note(tmp_path):
    import shutil

    copy = tmp_path / "lineage.db"
    shutil.copy(ROOT / "outputs" / "SP_CLO" / "lineage.db", copy)
    conn = sqlite3.connect(copy)
    obj = conn.execute("SELECT object_id FROM canonical_object WHERE object_name='CM_PAR_LOAN_TRADES'").fetchone()[0]
    conn.execute(
        "INSERT INTO canonical_column (column_id, object_id, column_name, declared_type, is_declared, is_inferred)"
        " VALUES ('COL_orphan_xl', ?, 'AAA_NO_LINEAGE', NULL, 0, 1)",
        (obj,),
    )
    conn.commit()
    conn.close()
    client = TestClient(create_app(str(copy)))
    object_id = _target_id(client, "CM_PAR_LOAN_TRADES")
    wb = _book(client.get(f"/targets/{object_id}/export.xlsx", params={"project": "SP_CLO"}))
    sheet = wb["Summary"]
    orphan = [row for row in sheet.iter_rows(min_row=2, values_only=True) if row[1] == "AAA_NO_LINEAGE"]
    assert len(orphan) == 1
    assert "No lineage" in orphan[0]


def test_legend_sheet_carries_traceability(sp_clo_client):
    object_id = _target_id(sp_clo_client, "CM_PAR_LOAN_TRADES")
    wb = _book(sp_clo_client.get(f"/targets/{object_id}/export.xlsx", params={"project": "SP_CLO"}))
    legend = [tuple(r) for r in wb["How to read"].iter_rows(values_only=True) if r[0]]
    assert ("Project", "SP_CLO") in legend
    assert any(r[0] == "Database" and r[1] == "lineage.db" for r in legend)
    assert any(r[0] == "Extracted at" and r[1] for r in legend)
    assert any(r[0] == "Lookup" for r in legend)
